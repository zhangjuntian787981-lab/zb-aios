import json
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt

from build_gm_approval_brief import (
    INK,
    LIGHT_GRAY,
    MUTED,
    add_bullet_item,
    add_callout,
    add_checkbox,
    add_form_line,
    add_metadata,
    add_numbering_definition,
    configure_styles,
    repeat_table_header,
    set_run_font,
    set_table_geometry,
    shade_cell,
)


ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / "implementation" / "p0" / "materials-and-recommendations.v1.json"
OUTPUT = ROOT / "outputs" / "中宝企业AI员工平台_P0资料清单与集中审批推荐包_v1.0.draft.docx"


def set_header_footer(section):
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header.paragraphs[0]
    header.alignment = WD_ALIGN_PARAGRAPH.LEFT
    run = header.add_run("P0 集中审批材料  |  中宝企业 AI 员工平台")
    set_run_font(run, size=8.5, color=MUTED, bold=True)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = footer.add_run("内部审批材料  ·  v1.0")
    set_run_font(run, size=8.5, color=MUTED)


def add_material_table(doc, rows, quantity_key):
    table = doc.add_table(rows=1, cols=3)
    widths = [2200, 4560, 2600]
    set_table_geometry(table, widths)
    for index, value in enumerate(("资料", "最少需要填写什么", "数量 / 处理方式")):
        cell = table.rows[0].cells[index]
        shade_cell(cell, LIGHT_GRAY)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(value)
        set_run_font(run, size=9.5, color=INK, bold=True)
    repeat_table_header(table.rows[0])

    for item in rows:
        quantity = item.get(quantity_key, "")
        if "safe_handling" in item:
            quantity = f"{quantity}\n{item['safe_handling']}"
        cells = table.add_row().cells
        values = (
            item["name"],
            "；".join(item.get("required_fields", item.get("must_include", []))),
            quantity,
        )
        for index, value in enumerate(values):
            paragraph = cells[index].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            run = paragraph.add_run(value)
            set_run_font(run, size=9.2, bold=index == 0)
        set_table_geometry(table, widths)
    return table


def add_recommendation_table(doc, recommendations):
    table = doc.add_table(rows=1, cols=3)
    widths = [1750, 4860, 2750]
    set_table_geometry(table, widths)
    for index, value in enumerate(("项目", "AI 推荐默认值", "Owner 只需核实")):
        cell = table.rows[0].cells[index]
        shade_cell(cell, LIGHT_GRAY)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(value)
        set_run_font(run, size=9.5, color=INK, bold=True)
    repeat_table_header(table.rows[0])

    for item in recommendations:
        cells = table.add_row().cells
        values = (item["title"], item["default"], item["owner_ack"])
        for index, value in enumerate(values):
            paragraph = cells[index].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            run = paragraph.add_run(value)
            set_run_font(run, size=9.1, bold=index == 0)
        set_table_geometry(table, widths)
    return table


def add_owner_ack_table(doc):
    rows = [
        ("业务 Owner", "真实流程、3—5 个任务、人工基线和可验收目标"),
        ("资料 Owner", "资料权威性、版本、有效期、权限和可进入 AI 范围"),
        ("IT / 身份 Owner", "登录、MFA、账号停用、撤权、测试环境和可复用能力"),
        ("安全 / 隐私 Owner", "数据等级、云模型边界、保留删除、审计和零容忍项"),
    ]
    table = doc.add_table(rows=1, cols=4)
    widths = [1700, 3700, 1980, 1980]
    set_table_geometry(table, widths)
    for index, value in enumerate(("责任人", "只核实这些事实", "姓名 / 代号", "结论 / 日期")):
        cell = table.rows[0].cells[index]
        shade_cell(cell, LIGHT_GRAY)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(value)
        set_run_font(run, size=9.2, color=INK, bold=True)
    repeat_table_header(table.rows[0])

    for role, scope in rows:
        cells = table.add_row().cells
        values = (role, scope, "\n\n", "无异议 / 有修订\n日期：")
        for index, value in enumerate(values):
            paragraph = cells[index].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            run = paragraph.add_run(value)
            set_run_font(run, size=9.1, bold=index == 0)
        set_table_geometry(table, widths)
    return table


def build_document():
    data = json.loads(SOURCE.read_text(encoding="utf-8"))
    doc = Document()
    doc.core_properties.title = "中宝企业 AI 员工平台：P0 资料清单与集中审批推荐包"
    doc.core_properties.subject = "P0 分阶段资料清单、AI 推荐默认值与集中审批"
    doc.core_properties.author = "中宝企业 AI 项目提议人"
    doc.core_properties.comments = "公司尚未授权；批准前仅为推荐草案。"

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    set_header_footer(section)
    configure_styles(doc)
    bullet_num = add_numbering_definition(doc, "bullet", "•")

    kicker = doc.add_paragraph("集中审批包  |  AI 推荐，待授权后签批", style="Memo Kicker")
    kicker.paragraph_format.space_before = Pt(10)
    doc.add_paragraph("P0 资料清单与集中审批推荐包", style="Memo Title")
    doc.add_paragraph(
        "你只补最少资料；AI 先给推荐；责任人只核实事实；最后由获授权的人一次性审批",
        style="Memo Subtitle",
    )
    add_metadata(doc, "项目", data["project"])
    add_metadata(doc, "包编号", data["package_id"])
    add_metadata(doc, "当前状态", "公司尚未授权；本包只是一份可修改的 AI 推荐草案")
    add_metadata(doc, "适用阶段", "P0 立项准备至 P1 Go / No-Go 之前")

    add_callout(
        doc,
        "推荐做法",
        "总经理先书面指定 P0 集中审批人；业务、资料、IT 和安全 Owner 只对本领域事实负责。AI 根据核实结果形成冻结版本和完整内容哈希，由集中审批人一次性批准、排除或退回。",
    )
    add_callout(
        doc,
        "不能被集中审批替代",
        "总经理的立项、人员、工时和预算授权；各 Owner 对真实事实和控制能力的证明；技术测试结果；以及 P1 最终 Go / No-Go。",
        tone="gray",
    )

    doc.add_heading("一、整个流程只有六步", level=1)
    for index, step in enumerate(data["approval_workflow"]["steps"], start=1):
        paragraph = doc.add_paragraph()
        paragraph.style = doc.styles["Normal"]
        number = paragraph.add_run(f"{index}. ")
        set_run_font(number, size=11, color=INK, bold=True)
        run = paragraph.add_run(step)
        set_run_font(run, size=11)

    doc.add_heading("二、你现在只需要提供五类资料", level=1)
    doc.add_paragraph(
        "现在不要提前索取公司敏感正文。没有的信息可以写“不知道”或使用推荐默认值。"
    )
    add_material_table(doc, data["materials_needed_now"], "minimum")

    doc.add_heading("三、总经理批准 P0 后再收集", level=1)
    doc.add_paragraph(
        "以下资料必须等总经理批准并任命负责人后再收集；当前不需要你一个人代替公司填写。"
    )
    add_material_table(doc, data["materials_after_p0_authorization"], "recommended_quantity")

    doc.add_heading("四、进入 P1 前必须补齐", level=1)
    doc.add_paragraph(
        "这些是小范围真实试点的放行证据。缺一项，阶段门保持 NOT_READY。"
    )
    add_material_table(doc, data["materials_required_before_p1"], "minimum")

    doc.add_page_break()
    doc.add_paragraph("AI 推荐默认值", style="Memo Title")
    doc.add_paragraph(
        "除非 Owner 提供相反事实或公司制度更严格，否则建议按下表形成初版",
        style="Memo Subtitle",
    )
    add_recommendation_table(doc, data["recommendations"])

    doc.add_heading("五、绝对不要提交的资料", level=1)
    for item in data["prohibited_materials"]:
        add_bullet_item(doc, bullet_num, item)

    doc.add_heading("六、Owner 事实核实表", level=1)
    doc.add_paragraph(
        "Owner 不需要审批整套技术方案，只需确认本行事实是否真实、是否有修订，并接受本领域后续维护责任。"
    )
    add_owner_ack_table(doc)
    add_form_line(doc, "Owner 修订项汇总 / 证据位置", width=54)

    doc.add_page_break()
    doc.add_paragraph("P0 集中审批单", style="Memo Title")
    doc.add_paragraph(
        "仅限总经理书面授权的 P0 集中审批人填写；未授权、Owner 未核实或没有内容哈希均不生效",
        style="Memo Subtitle",
    )

    doc.add_heading("A. 授权与前置事实", level=1)
    add_form_line(doc, "总经理授权记录编号 / 位置", width=38)
    add_form_line(doc, "P0 集中审批人姓名 / 代号", width=34)
    add_form_line(doc, "P1 Go / No-Go 审批人姓名 / 代号", width=30)
    add_form_line(doc, "四类 Owner 核实记录位置", width=38)

    doc.add_heading("B. 本次集中审批决定", level=1)
    add_checkbox(doc, "全部批准：同意按冻结版本进入 P0 后续准备")
    add_checkbox(doc, "部分批准：除下列排除项外，其余内容批准")
    add_checkbox(doc, "退回修改：Owner 事实或推荐方案需要重新核实")
    add_checkbox(doc, "暂缓：等待公司资源、政策或其他前置条件")
    add_form_line(doc, "排除项 / 退回原因", width=48)

    doc.add_heading("C. 审批对象必须唯一可复查", level=1)
    add_form_line(doc, "推荐包版本", width=28)
    add_form_line(doc, "冻结文件位置", width=42)
    add_form_line(doc, "完整内容 SHA-256 哈希", width=38)
    add_form_line(doc, "审批证据位置", width=42)

    doc.add_heading("D. 签字", level=1)
    add_form_line(doc, "P0 集中审批人签字", width=30)
    add_form_line(doc, "审批日期和时间", width=30)
    add_form_line(doc, "下一次复审日期", width=30)

    add_callout(
        doc,
        "审批后仍然不能做",
        "本次集中审批不授权接入 OA、U9、BI，不授权真实系统写回、自动对外发送、员工 KPI 分析或进入 P1。进入 P1 必须由签批单指定的阶段门审批人另行作出 Go / No-Go。",
    )

    doc.add_heading("E. P1 最终阶段门", level=1)
    add_checkbox(doc, "GO：全部公司事实、技术证据和集中审批均有效，可按批准范围进入 P1")
    add_checkbox(doc, "NO-GO：停止进入 P1，并记录原因和整改条件")
    add_form_line(doc, "P1 阶段门审批人签字", width=26)
    add_form_line(doc, "决定日期和证据位置", width=30)

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    raise SystemExit(
        "Archived generator: use docs/plans and implementation/p0 from the v4 product model."
    )
"""Archived v3 company-internal approval pack generator.

The active v4 product model uses an external Product Owner, synthetic data
through P2, and enterprise onboarding in P3. This file is retained only to
preserve the historical artifact recipe.
"""
