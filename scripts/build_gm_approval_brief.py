from pathlib import Path

from docx import Document
from docx.enum.table import WD_ALIGN_VERTICAL
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
OUTPUT = ROOT / "outputs" / "中宝企业AI员工平台_总经理首次审批说明与签批单_v1.0.draft.docx"

PAGE_WIDTH_DXA = 9360
TABLE_INDENT_DXA = 120
BLUE = "2E74B5"
DARK_BLUE = "1F4D78"
INK = "0B2545"
MUTED = "5B6573"
LIGHT_BLUE = "EAF2F8"
LIGHT_GRAY = "F2F4F7"
BORDER = "C7D0D9"
RED = "9B1C1C"
GOLD = "7A5A00"


def set_run_font(run, size=None, color="000000", bold=None, italic=None):
    run.font.name = "Arial"
    run._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Hiragino Sans GB")
    if size is not None:
        run.font.size = Pt(size)
    run.font.color.rgb = RGBColor.from_string(color)
    if bold is not None:
        run.bold = bold
    if italic is not None:
        run.italic = italic


def set_style_font(style, font_size, color, bold=False):
    style.font.name = "Arial"
    style._element.get_or_add_rPr().rFonts.set(qn("w:ascii"), "Arial")
    style._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"), "Arial")
    style._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), "Hiragino Sans GB")
    style.font.size = Pt(font_size)
    style.font.color.rgb = RGBColor.from_string(color)
    style.font.bold = bold


def configure_styles(doc):
    normal = doc.styles["Normal"]
    set_style_font(normal, 11, "000000")
    normal.paragraph_format.space_before = Pt(0)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.1
    normal.paragraph_format.widow_control = True

    h1 = doc.styles["Heading 1"]
    set_style_font(h1, 16, BLUE, True)
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    h1.paragraph_format.keep_with_next = True
    h1.paragraph_format.keep_together = True

    h2 = doc.styles["Heading 2"]
    set_style_font(h2, 13, BLUE, True)
    h2.paragraph_format.space_before = Pt(10)
    h2.paragraph_format.space_after = Pt(5)
    h2.paragraph_format.keep_with_next = True
    h2.paragraph_format.keep_together = True

    h3 = doc.styles["Heading 3"]
    set_style_font(h3, 12, DARK_BLUE, True)
    h3.paragraph_format.space_before = Pt(8)
    h3.paragraph_format.space_after = Pt(4)
    h3.paragraph_format.keep_with_next = True

    title = doc.styles.add_style("Memo Title", 1)
    set_style_font(title, 23, "000000", True)
    title.paragraph_format.space_before = Pt(0)
    title.paragraph_format.space_after = Pt(4)
    title.paragraph_format.line_spacing = 1.0
    title.paragraph_format.keep_with_next = True

    subtitle = doc.styles.add_style("Memo Subtitle", 1)
    set_style_font(subtitle, 13, MUTED, False)
    subtitle.paragraph_format.space_before = Pt(0)
    subtitle.paragraph_format.space_after = Pt(14)
    subtitle.paragraph_format.line_spacing = 1.1
    subtitle.paragraph_format.keep_with_next = True

    kicker = doc.styles.add_style("Memo Kicker", 1)
    set_style_font(kicker, 10, BLUE, True)
    kicker.paragraph_format.space_before = Pt(0)
    kicker.paragraph_format.space_after = Pt(6)
    kicker.paragraph_format.keep_with_next = True

    muted = doc.styles.add_style("Small Muted", 1)
    set_style_font(muted, 9, MUTED, False)
    muted.paragraph_format.space_before = Pt(0)
    muted.paragraph_format.space_after = Pt(4)
    muted.paragraph_format.line_spacing = 1.0

    form = doc.styles.add_style("Form Prompt", 1)
    set_style_font(form, 10.5, "000000", False)
    form.paragraph_format.space_before = Pt(2)
    form.paragraph_format.space_after = Pt(7)
    form.paragraph_format.line_spacing = 1.1


def set_cell_margins(cell, top=100, start=120, bottom=100, end=120):
    tc = cell._tc
    tc_pr = tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for margin, value in (
        ("top", top),
        ("start", start),
        ("bottom", bottom),
        ("end", end),
    ):
        node = tc_mar.find(qn(f"w:{margin}"))
        if node is None:
            node = OxmlElement(f"w:{margin}")
            tc_mar.append(node)
        node.set(qn("w:w"), str(value))
        node.set(qn("w:type"), "dxa")


def set_cell_width(cell, width_dxa):
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_w = tc_pr.first_child_found_in("w:tcW")
    if tc_w is None:
        tc_w = OxmlElement("w:tcW")
        tc_pr.append(tc_w)
    tc_w.set(qn("w:w"), str(width_dxa))
    tc_w.set(qn("w:type"), "dxa")


def set_table_geometry(table, widths):
    table.autofit = False
    tbl = table._tbl
    tbl_pr = tbl.tblPr

    tbl_w = tbl_pr.first_child_found_in("w:tblW")
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_w.set(qn("w:type"), "dxa")

    tbl_ind = tbl_pr.first_child_found_in("w:tblInd")
    if tbl_ind is None:
        tbl_ind = OxmlElement("w:tblInd")
        tbl_pr.append(tbl_ind)
    tbl_ind.set(qn("w:w"), str(TABLE_INDENT_DXA))
    tbl_ind.set(qn("w:type"), "dxa")

    layout = tbl_pr.first_child_found_in("w:tblLayout")
    if layout is None:
        layout = OxmlElement("w:tblLayout")
        tbl_pr.append(layout)
    layout.set(qn("w:type"), "fixed")

    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        border = borders.find(qn(f"w:{edge}"))
        if border is None:
            border = OxmlElement(f"w:{edge}")
            borders.append(border)
        border.set(qn("w:val"), "single")
        border.set(qn("w:sz"), "6")
        border.set(qn("w:color"), BORDER)

    grid = tbl.tblGrid
    for child in list(grid):
        grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        grid.append(col)

    for row in table.rows:
        for index, cell in enumerate(row.cells):
            set_cell_width(cell, widths[index])
            set_cell_margins(cell)
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER


def shade_cell(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.first_child_found_in("w:shd")
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def repeat_table_header(row):
    tr_pr = row._tr.get_or_add_trPr()
    header = OxmlElement("w:tblHeader")
    header.set(qn("w:val"), "true")
    tr_pr.append(header)


def add_paragraph_border(paragraph, side, color, size=14, space=6):
    p_pr = paragraph._p.get_or_add_pPr()
    p_bdr = p_pr.first_child_found_in("w:pBdr")
    if p_bdr is None:
        p_bdr = OxmlElement("w:pBdr")
        p_pr.append(p_bdr)
    edge = OxmlElement(f"w:{side}")
    edge.set(qn("w:val"), "single")
    edge.set(qn("w:sz"), str(size))
    edge.set(qn("w:space"), str(space))
    edge.set(qn("w:color"), color)
    p_bdr.append(edge)


def add_callout(doc, label, text, tone="blue"):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(6)
    paragraph.paragraph_format.space_after = Pt(8)
    paragraph.paragraph_format.left_indent = Inches(0.12)
    paragraph.paragraph_format.right_indent = Inches(0.08)
    paragraph.paragraph_format.line_spacing = 1.12
    p_pr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:fill"), LIGHT_BLUE if tone == "blue" else LIGHT_GRAY)
    p_pr.append(shd)
    add_paragraph_border(
        paragraph,
        "left",
        BLUE if tone == "blue" else GOLD,
        size=22,
        space=8,
    )
    label_run = paragraph.add_run(f"{label}：")
    set_run_font(label_run, size=11, color=INK, bold=True)
    text_run = paragraph.add_run(text)
    set_run_font(text_run, size=11, color="000000")
    return paragraph


def add_numbering_definition(doc, num_format, text, left=720, hanging=360):
    numbering = doc.part.numbering_part.element
    abstract_ids = [
        int(node.get(qn("w:abstractNumId")))
        for node in numbering.findall(qn("w:abstractNum"))
    ]
    num_ids = [
        int(node.get(qn("w:numId"))) for node in numbering.findall(qn("w:num"))
    ]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1

    abstract = OxmlElement("w:abstractNum")
    abstract.set(qn("w:abstractNumId"), str(abstract_id))
    multi = OxmlElement("w:multiLevelType")
    multi.set(qn("w:val"), "singleLevel")
    abstract.append(multi)

    level = OxmlElement("w:lvl")
    level.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start")
    start.set(qn("w:val"), "1")
    level.append(start)
    fmt = OxmlElement("w:numFmt")
    fmt.set(qn("w:val"), num_format)
    level.append(fmt)
    level_text = OxmlElement("w:lvlText")
    level_text.set(qn("w:val"), text)
    level.append(level_text)
    level_jc = OxmlElement("w:lvlJc")
    level_jc.set(qn("w:val"), "left")
    level.append(level_jc)

    p_pr = OxmlElement("w:pPr")
    tabs = OxmlElement("w:tabs")
    tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num")
    tab.set(qn("w:pos"), str(left))
    tabs.append(tab)
    p_pr.append(tabs)
    ind = OxmlElement("w:ind")
    ind.set(qn("w:left"), str(left))
    ind.set(qn("w:hanging"), str(hanging))
    p_pr.append(ind)
    spacing = OxmlElement("w:spacing")
    spacing.set(qn("w:after"), "160")
    spacing.set(qn("w:line"), "280")
    spacing.set(qn("w:lineRule"), "auto")
    p_pr.append(spacing)
    level.append(p_pr)
    abstract.append(level)
    numbering.append(abstract)

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(num_id))
    abstract_num_id = OxmlElement("w:abstractNumId")
    abstract_num_id.set(qn("w:val"), str(abstract_id))
    num.append(abstract_num_id)
    numbering.append(num)
    return num_id


def add_numbered_item(doc, num_id, text, bold_lead=None):
    paragraph = doc.add_paragraph()
    p_pr = paragraph._p.get_or_add_pPr()
    num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl")
    ilvl.set(qn("w:val"), "0")
    num_id_node = OxmlElement("w:numId")
    num_id_node.set(qn("w:val"), str(num_id))
    num_pr.append(ilvl)
    num_pr.append(num_id_node)
    p_pr.append(num_pr)
    if bold_lead and text.startswith(bold_lead):
        lead = paragraph.add_run(bold_lead)
        set_run_font(lead, size=11, bold=True)
        body = paragraph.add_run(text[len(bold_lead) :])
        set_run_font(body, size=11)
    else:
        run = paragraph.add_run(text)
        set_run_font(run, size=11)
    return paragraph


def add_bullet_item(doc, num_id, text):
    return add_numbered_item(doc, num_id, text)


def add_metadata(doc, label, value):
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0)
    paragraph.paragraph_format.space_after = Pt(2)
    paragraph.paragraph_format.line_spacing = 1.0
    label_run = paragraph.add_run(f"{label}：")
    set_run_font(label_run, size=10.5, bold=True)
    value_run = paragraph.add_run(value)
    set_run_font(value_run, size=10.5)


def add_form_line(doc, label, width=36):
    paragraph = doc.add_paragraph(style="Form Prompt")
    label_run = paragraph.add_run(f"{label}：")
    set_run_font(label_run, size=10.5, bold=True)
    value_run = paragraph.add_run("_" * width)
    set_run_font(value_run, size=10.5)
    return paragraph


def add_checkbox(doc, text):
    paragraph = doc.add_paragraph(style="Form Prompt")
    marker = paragraph.add_run("☐ ")
    set_run_font(marker, size=12)
    run = paragraph.add_run(text)
    set_run_font(run, size=10.5)
    return paragraph


def set_header_footer(section):
    section.header_distance = Inches(0.492)
    section.footer_distance = Inches(0.492)

    header = section.header
    header_para = header.paragraphs[0]
    header_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    header_para.paragraph_format.space_after = Pt(0)
    header_run = header_para.add_run("总经理决策简报  |  中宝企业 AI 员工平台")
    set_run_font(header_run, size=8.5, color=MUTED, bold=True)

    footer = section.footer
    footer_para = footer.paragraphs[0]
    footer_para.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    footer_para.paragraph_format.space_before = Pt(0)
    label = footer_para.add_run("内部审批材料  ·  v1.0")
    set_run_font(label, size=8.5, color=MUTED)


def add_role_table(doc):
    rows = [
        ("执行赞助人/最终决策人", "总经理本人或书面授权的分管领导", "决定资源、处理跨部门分歧、作阶段 Go/No-Go"),
        ("项目负责人", "建议考虑当前项目提议人；由总经理决定", "组织计划、协调、记录、风险上报；不能独自验收自己"),
        ("业务负责人", "首批候选部门安排 1 人", "定义真实问题、提供案例、判断是否有业务价值"),
        ("资料负责人", "相关资料部门安排 1—2 人", "批准资料、版本、查看范围、更新和撤回"),
        ("IT/身份负责人", "现有信息化人员或受控支持人员", "确认登录、账号停用、权限、备份和故障处理"),
        ("安全/隐私负责人", "安全、合规或管理层指定人员", "批准敏感数据边界、云端规则和事故处理"),
    ]
    table = doc.add_table(rows=1, cols=3)
    set_table_geometry(table, [2100, 3000, 4260])
    headers = ["角色", "总经理需要任命谁", "承担什么责任"]
    for index, value in enumerate(headers):
        cell = table.rows[0].cells[index]
        shade_cell(cell, LIGHT_GRAY)
        paragraph = cell.paragraphs[0]
        paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
        run = paragraph.add_run(value)
        set_run_font(run, size=9.5, color=INK, bold=True)
    repeat_table_header(table.rows[0])

    for role, assignee, responsibility in rows:
        cells = table.add_row().cells
        for index, value in enumerate((role, assignee, responsibility)):
            paragraph = cells[index].paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(0)
            paragraph.paragraph_format.line_spacing = 1.05
            run = paragraph.add_run(value)
            set_run_font(run, size=9.5, bold=index == 0)
        set_table_geometry(table, [2100, 3000, 4260])
    return table


def build_document():
    doc = Document()
    doc.core_properties.title = "中宝企业 AI 员工平台：总经理首次审批说明与签批单"
    doc.core_properties.subject = "P0 立项申请与人员任命"
    doc.core_properties.author = "中宝企业 AI 项目提议人"
    doc.core_properties.comments = "公司尚未立项；本文件为待审批材料。"

    section = doc.sections[0]
    section.page_width = Inches(8.5)
    section.page_height = Inches(11)
    section.top_margin = Inches(1)
    section.bottom_margin = Inches(1)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)
    set_header_footer(section)
    configure_styles(doc)

    decimal_num = add_numbering_definition(doc, "decimal", "%1.")
    bullet_num = add_numbering_definition(doc, "bullet", "•")

    kicker = doc.add_paragraph("决策备忘录  |  待总经理审批", style="Memo Kicker")
    kicker.paragraph_format.space_before = Pt(10)
    doc.add_paragraph(
        "关于启动“中宝企业 AI 员工平台”P0 准备阶段的请示",
        style="Memo Title",
    )
    doc.add_paragraph(
        "您现在批准的是 2—3 周的可行性准备，不是全公司上线，也不是后续全部预算",
        style="Memo Subtitle",
    )

    add_metadata(doc, "呈报对象", "总经理")
    add_metadata(doc, "呈报人", "项目提议人（当前唯一执行者，尚未正式任命）")
    add_metadata(doc, "日期", "2026 年 7 月 25 日")
    add_metadata(doc, "当前状态", "公司尚未立项；仅完成部分非生产方案、测试和进度工具")

    add_callout(
        doc,
        "建议决定",
        "原则同意启动一个有边界、有负责人、有截止日期的 P0；任命责任人后再开展跨部门准备，P0 完成后重新审批是否进入 P1 小范围试点。",
    )

    doc.add_heading("一、这是什么项目", level=1)
    paragraph = doc.add_paragraph()
    paragraph.add_run("一句话解释：").bold = True
    paragraph.add_run(
        "建设一个由公司统一管理的员工 AI 入口。每名员工有自己的账号和个人助手，公司决定其能看到哪些资料、能使用哪些工作方法、哪些操作必须审批。"
    )
    doc.add_paragraph(
        "未来条件成熟后，可以通过正式接口连接 OA、U9、BI；但本次申请不连接任何生产系统，也不申请全公司推广。"
    )

    doc.add_heading("二、为什么现在必须由总经理决定", level=1)
    doc.add_paragraph(
        "目前只有项目提议人一人在准备方案、测试和进度工具。这些成果可以证明“有人做过准备”，但不能代表公司已经立项，也不能由一个人代替各部门决定业务价值、资料权限、员工参与、隐私边界和验收结果。"
    )
    add_callout(
        doc,
        "治理底线",
        "项目负责人不能同时成为唯一的业务验收人、唯一的资料批准人和唯一的安全批准人。未经总经理书面授权，不调动员工、不索取公司资料、不进入 P1。",
        tone="gray",
    )

    doc.add_heading("三、您现在只需要作出六项决定", level=1)
    decisions = [
        ("是否批准启动 P0。", "只批准 2—3 周的准备工作，不一次批准 P1—P3、全员上线或生产系统接入。"),
        ("由谁作最终决策。", "由您本人或书面授权的一名分管领导担任执行赞助人，协调部门并处理分歧。"),
        ("正式任命项目负责人。", "可考虑由当前提议人担任，也可以另行指定；必须给出明确工作时间和汇报要求。"),
        ("安排跨部门责任人。", "业务、资料、IT/身份、安全/隐私分别有人承担责任，不能只写“相关部门”。"),
        ("只选一个首批主任务。", "建议从低风险、可人工核对的内部知识查询开始，最多另设一个备用任务。"),
        ("确认投入、红线和复审。", "明确工时、费用上限、P0 截止日期，以及 P0 结束后的再次审批日期。"),
    ]
    for lead, detail in decisions:
        add_numbered_item(doc, decimal_num, f"{lead}{detail}", bold_lead=lead)

    doc.add_heading("四、需要任命的最小团队", level=1)
    doc.add_paragraph(
        "建议由 4—6 名兼职人员承担以下六类职责，不需要现在新建部门，也不需要发动全公司。部分角色可兼任，但项目执行、业务验收、资料批准和安全批准不能全部由同一人完成。"
    )
    add_role_table(doc)

    doc.add_heading("五、P0 明确不做什么", level=1)
    boundaries = [
        "不连接 OA、U9、BI 或其他生产系统，不使用共享账号、数据库账号、Cookie 或网页抓取绕过限制。",
        "不自动发送邮件、客户回复、报价、合同或任何对外承诺。",
        "不修改订单、审批、财务、采购、生产或人事数据。",
        "不开放员工任意代码执行或不受控的互联网访问。",
        "不分析个人员工 KPI，不用于员工排名、奖惩、晋升、辞退等决定。",
        "不把模拟结果、测试通过或准备材料宣传为已经生产上线。",
        "不一次性采购大规模服务器；任何费用超过本次批准上限必须另报。",
    ]
    for item in boundaries:
        add_bullet_item(doc, bullet_num, item)

    doc.add_heading("六、P0 结束后您将收到什么", level=1)
    deliverables = [
        "一个主试点任务、业务负责人和可人工核对的成功标准；",
        "经负责人批准的资料清单、版本、权限和撤回规则；",
        "候选试点人员、登录与离职撤权建议；",
        "数据、隐私、权限和安全红线；",
        "人工现状基线、风险、费用和人员投入说明；",
        "“进入 P1 / 补充整改 / 停止项目”的书面建议。",
    ]
    for item in deliverables:
        add_bullet_item(doc, bullet_num, item)
    add_callout(
        doc,
        "第二道审批门",
        "P0 获批不等于 P1 获批。只有总经理或书面授权人再次签批，才可以安排小范围试点；试点人数需届时单独确认。",
    )

    doc.add_heading("七、总经理不需要亲自做什么", level=1)
    doc.add_paragraph(
        "您不需要选择 AI 模型、服务器、编程语言或开源软件，也不需要审核每一份技术文档。技术选型和测试由正式任命的项目负责人组织，业务内容和资料分别由业务负责人、资料负责人把关。"
    )
    doc.add_paragraph(
        "您需要做的是：决定公司是否值得用少量时间先验证、由谁负责、给多少资源、哪些事绝对不能做，以及什么时候回来复审。"
    )

    doc.add_page_break()
    doc.add_paragraph("总经理签批单", style="Memo Title")
    doc.add_paragraph(
        "本次只审批 P0 可行性与试点准备；空白、口头同意或未填写责任人均不视为通过。",
        style="Memo Subtitle",
    )

    doc.add_heading("A. 本次决定", level=1)
    add_checkbox(doc, "同意启动 P0 准备阶段（建议项），完成后重新审批是否进入 P1")
    add_checkbox(doc, "有条件同意，条件为：________________________________________")
    add_checkbox(doc, "修改后再审，需补充：________________________________________")
    add_checkbox(doc, "暂缓或不同意，原因：________________________________________")

    doc.add_heading("B. 人员任命", level=1)
    add_form_line(doc, "执行赞助人 / 最终决策人")
    add_form_line(doc, "项目负责人")
    add_form_line(doc, "首批候选部门")
    add_form_line(doc, "业务负责人")
    add_form_line(doc, "资料负责人")
    add_form_line(doc, "IT / 身份负责人")
    add_form_line(doc, "安全 / 隐私负责人")

    doc.add_heading("C. 首批范围与资源", level=1)
    add_form_line(doc, "唯一主任务")
    add_form_line(doc, "备用任务（可空）")
    add_form_line(doc, "P0 周期", width=26)
    add_form_line(doc, "项目负责人每周可投入时间", width=24)
    add_checkbox(doc, "P0 原则上不新增采购，仅使用现有资源")
    add_form_line(doc, "允许的验证费用上限（如有）", width=20)

    doc.add_heading("D. 必须遵守的边界", level=1)
    approvals = [
        "不连接 OA、U9、BI 等生产系统",
        "不使用生产账号、密码、Cookie 和未批准敏感数据",
        "不自动对外发送，不作价格、交期、参数、认证或质保承诺",
        "不写入任何业务系统，不开放员工任意代码执行",
        "不分析个人员工 KPI，不用于员工奖惩",
        "任何扩围、采购、系统接入和 P1 试点均重新审批",
    ]
    for item in approvals:
        add_checkbox(doc, item)

    doc.add_heading("E. 复审与签字", level=1)
    add_form_line(doc, "P0 结果汇报日期", width=22)
    add_form_line(doc, "总经理 / 授权人签字", width=22)
    add_form_line(doc, "签批日期", width=22)

    add_callout(
        doc,
        "推荐批示",
        "同意启动 P0 可行性与试点准备。本次批准不代表全公司上线，不代表进入 P1，也不授权连接 OA、U9、BI 或其他生产系统。未经再次书面批准，不得扩大人员、增加敏感数据、分析个人员工 KPI、对外发送或执行正式业务写入。",
    )

    doc.save(OUTPUT)
    print(OUTPUT)


if __name__ == "__main__":
    build_document()
